#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]

from openpyxl import load_workbook

from audits.jake_audit_workbook import DEFAULT_TEMPLATE_WORKBOOK, generate_nycha_audit_workbook
from mcp.jake_ops_mcp import JakeOps


def donor_addresses(template_path: Path) -> list[str]:
    wb = load_workbook(template_path, read_only=True)
    return [name for name in wb.sheetnames if str(name).strip().lower() != "template"]


def markdown_table(rows: list[dict[str, object]]) -> str:
    lines = [
        "| Address | Ready % | Good | Seen/Wrong | No Evidence | Layout | Workbook |",
        "| --- | ---: | ---: | ---: | ---: | --- | --- |",
    ]
    for row in rows:
        lines.append(
            f"| {row['address']} | {row['weighted_ready_percent']} | {row['good_count']} | "
            f"{row['seen_wrong_count']} | {row['no_evidence_count']} | {row['layout_kind']} | {row['output_path']} |"
        )
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate batch NYCHA audit readiness summary from Jake.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE_WORKBOOK), help="Donor workbook path used for address/sheet list")
    parser.add_argument("--threshold", type=int, default=70, help="Readiness threshold to flag")
    parser.add_argument("--limit", type=int, default=0, help="Optional max number of addresses to process")
    parser.add_argument(
        "--out-dir",
        default=str(PROJECT_ROOT / "output" / "spreadsheet"),
        help="Workbook output directory",
    )
    args = parser.parse_args()

    template_path = Path(args.template)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    ops = JakeOps()
    summary_rows: list[dict[str, object]] = []
    errors: list[dict[str, str]] = []

    addresses = donor_addresses(template_path)
    if args.limit and args.limit > 0:
        addresses = addresses[: args.limit]

    for index, address in enumerate(addresses, start=1):
        print(f"[{index}/{len(addresses)}] auditing {address}...", file=sys.stderr, flush=True)
        try:
            output_path = out_dir / f"{address}_audit.xlsx"
            result = generate_nycha_audit_workbook(
                address_text=address,
                out_path=output_path,
                template_path=template_path,
                ops=ops,
            )
            rows = result.get("rows") or []
            good_count = sum(1 for row in rows if row.get("state") == "green")
            seen_wrong_count = sum(1 for row in rows if row.get("state") == "yellow")
            no_evidence_count = sum(1 for row in rows if row.get("state") == "red")
            summary_rows.append(
                {
                    "address": address,
                    "weighted_ready_percent": int(result.get("weighted_ready_percent") or 0),
                    "good_count": good_count,
                    "seen_wrong_count": seen_wrong_count,
                    "no_evidence_count": no_evidence_count,
                    "layout_kind": str(result.get("layout_kind") or ""),
                    "output_path": str(result.get("output_path") or output_path),
                    "live_site_id": str(result.get("live_site_id") or ""),
                    "live_building_id": str(result.get("live_building_id") or ""),
                    "live_alert_count": int(result.get("live_alert_count") or 0),
                    "row_count": int(result.get("row_count") or 0),
                }
            )
            print(
                f"[{index}/{len(addresses)}] done {address}: {int(result.get('weighted_ready_percent') or 0)}%",
                file=sys.stderr,
                flush=True,
            )
        except Exception as exc:
            errors.append({"address": address, "error": str(exc)})
            print(f"[{index}/{len(addresses)}] error {address}: {exc}", file=sys.stderr, flush=True)

    summary_rows.sort(key=lambda row: (int(row["weighted_ready_percent"]), str(row["address"])))
    below_threshold = [row for row in summary_rows if int(row["weighted_ready_percent"]) < args.threshold]

    csv_path = PROJECT_ROOT / "output" / "spreadsheet" / "nycha_audit_readiness_summary.csv"
    json_path = PROJECT_ROOT / "output" / "spreadsheet" / "nycha_audit_readiness_summary.json"
    md_path = PROJECT_ROOT / "output" / "spreadsheet" / "nycha_audit_readiness_summary.md"

    with csv_path.open("w", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "address",
                "weighted_ready_percent",
                "good_count",
                "seen_wrong_count",
                "no_evidence_count",
                "layout_kind",
                "live_site_id",
                "live_building_id",
                "live_alert_count",
                "row_count",
                "output_path",
            ],
        )
        writer.writeheader()
        writer.writerows(summary_rows)

    json_path.write_text(
        json.dumps(
            {
                "threshold": args.threshold,
                "site_count": len(summary_rows),
                "below_threshold_count": len(below_threshold),
                "below_threshold": below_threshold,
                "all_sites": summary_rows,
                "errors": errors,
            },
            indent=2,
        )
    )

    md_lines = [
        f"# NYCHA Audit Readiness Summary",
        "",
        f"- Threshold: {args.threshold}%",
        f"- Total audited sites: {len(summary_rows)}",
        f"- Sites below threshold: {len(below_threshold)}",
        "",
        "## Below Threshold",
        "",
    ]
    if below_threshold:
        md_lines.append(markdown_table(below_threshold))
    else:
        md_lines.append("None.")
    md_lines.extend(["", "## All Sites", "", markdown_table(summary_rows)])
    if errors:
        md_lines.extend(["", "## Errors", ""])
        for row in errors:
            md_lines.append(f"- {row['address']}: {row['error']}")
    md_path.write_text("\n".join(md_lines))

    print(json.dumps({
        "threshold": args.threshold,
        "site_count": len(summary_rows),
        "below_threshold_count": len(below_threshold),
        "csv_path": str(csv_path),
        "json_path": str(json_path),
        "md_path": str(md_path),
        "errors": errors,
    }, indent=2))


if __name__ == "__main__":
    main()
